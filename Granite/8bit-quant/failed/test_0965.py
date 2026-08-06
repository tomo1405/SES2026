import os
import pandas as pd
import docx
import pytest
from pathlib import Path
from src_0965 import task_func

@pytest.fixture
def setup_directories(tmp_path):
    source_directory = tmp_path / "source"
    target_directory = tmp_path / "target"
    source_directory.mkdir()
    target_directory.mkdir()
    return str(source_directory), str(target_directory)

def test_task_func_valid_directories(setup_directories):
    source_directory, target_directory = setup_directories
    assert task_func(source_directory, target_directory) == 0

def test_task_func_invalid_source_directory(setup_directories):
    source_directory, target_directory = setup_directories
    invalid_directory = "/invalid/directory"
    with pytest.raises(FileNotFoundError):
        task_func(invalid_directory, target_directory)

def test_task_func_invalid_target_directory(setup_directories):
    source_directory, target_directory = setup_directories
    invalid_directory = "/invalid/directory"
    os.makedirs(invalid_directory, exist_ok=True)
    with pytest.raises(FileNotFoundError):
        task_func(source_directory, invalid_directory)

def test_task_func_convert_csv(setup_directories):
    source_directory, target_directory = setup_directories
    csv_file = source_directory + "/data.csv"
    df = pd.DataFrame({"A": [1, 2, 3], "B": [4, 5, 6]})
    df.to_csv(csv_file, index=False)
    assert task_func(source_directory, target_directory) == 1
    converted_csv_file = target_directory + "/data.csv"
    assert os.path.exists(converted_csv_file)
    converted_df = pd.read_csv(converted_csv_file)
    assert converted_df.equals(df)

def test_task_func_convert_excel(setup_directories):
    source_directory, target_directory = setup_directories
    excel_file = source_directory + "/data.xlsx"
    df = pd.DataFrame({"A": [1, 2, 3], "B": [4, 5, 6]})
    df.to_excel(excel_file, index=False)
    assert task_func(source_directory, target_directory) == 1
    converted_excel_file = target_directory + "/data.csv"
    assert os.path.exists(converted_excel_file)
    converted_df = pd.read_csv(converted_excel_file)
    assert converted_df.equals(df)

def test_task_func_convert_docx(setup_directories):
    source_directory, target_directory = setup_directories
    docx_file = source_directory + "/data.docx"
    doc = docx.Document()
    doc.add_paragraph("Hello, world!")
    doc.add_paragraph("This is a test.")
    doc.save(docx_file)
    assert task_func(source_directory, target_directory) == 1
    converted_docx_file = target_directory + "/data.csv"
    assert os.path.exists(converted_docx_file)
    converted_df = pd.read_csv(converted_docx_file)
    assert converted_df["Text"].iloc[0] == "Hello, world!"
    assert converted_df["Text"].iloc[1] == "This is a test."

def test_task_func_convert_txt(setup_directories):
    source_directory, target_directory = setup_directories
    txt_file = source_directory + "/data.txt"
    with open(txt_file, "w") as f:
        f.write("Hello, world!\n")
        f.write("This is a test.")
    assert task_func(source_directory, target_directory) == 1
    converted_txt_file = target_directory + "/data.csv"
    assert os.path.exists(converted_txt_file)
    converted_df = pd.read_csv(converted_txt_file)
    assert converted_df["Text"].iloc[0] == "Hello, world!"
    assert converted_df["Text"].iloc[1] == "This is a test."