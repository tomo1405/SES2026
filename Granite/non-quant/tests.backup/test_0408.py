import os
import csv
from openpyxl import load_workbook
from unittest.mock import patch, call

def task_func(file_name, excel_file_path, csv_file_path) -> str:

    excel_file = os.path.join(excel_file_path, file_name)
    # Check if the Excel file exists
    if not os.path.isfile(excel_file):
        raise FileNotFoundError(f"[Errno 2] No such file or directory: '{excel_file}'")

    workbook = load_workbook(filename=excel_file, read_only=True)
    sheet = workbook.active

    data = [[cell.value for cell in row] for row in sheet.iter_rows()]

    csv_file_name = os.path.splitext(file_name)[0] + '.csv'
    csv_file = os.path.join(csv_file_path, csv_file_name)

    with open(csv_file, 'w', newline='', encoding='utf-8') as file:
        writer = csv.writer(file)
        writer.writerows(data)

    return csv_file_name

def test_task_func():
    with patch('os.path.isfile') as mock_isfile, \
         patch('openpyxl.load_workbook') as mock_load_workbook, \
         patch('csv.writer') as mock_writer:

        mock_isfile.return_value = True
        mock_workbook = mock_load_workbook.return_value
        mock_sheet = mock_workbook.active
        mock_data = [[cell.value for cell in row] for row in mock_sheet.iter_rows()]

        task_func('test_file.xlsx', 'excel_path', 'csv_path')

        mock_isfile.assert_called_once_with(os.path.join('excel_path', 'test_file.xlsx'))
        mock_load_workbook.assert_called_once_with(filename=os.path.join('excel_path', 'test_file.xlsx'), read_only=True)
        mock_sheet.iter_rows.assert_called_once()
        mock_writer.assert_called_once()
        mock_writer.return_value.writerows.assert_called_once_with(mock_data)