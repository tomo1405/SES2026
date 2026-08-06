import glob
import os

import pytest
from openpyxl import Workbook
from src_0375 import task_func


def test_task_func_no_directory():
    with pytest.raises(FileNotFoundError, match='The specified directory does not exist.'):
        task_func('./non_existent_directory/')

def test_task_func_empty_directory(tmpdir):
    result = task_func(str(tmpdir))
    assert result == 0

def test_task_func_with_xlsx_files(tmpdir):
    # Create some test .xlsx files
    wb1 = Workbook()
    ws1 = wb1.active
    ws1['A1'] = 'Test "String"'
    wb1.save(os.path.join(tmpdir, 'test1.xlsx'))

    wb2 = Workbook()
    ws2 = wb2.active
    ws2['B2'] = 'Another "Test" String'
    wb2.save(os.path.join(tmpdir, 'test2.xlsx'))

    result = task_func(str(tmpdir))
    assert result == 2

    # Check if the files were processed correctly
    for filename in ['test1.xlsx', 'test2.xlsx']:
        workbook = load_workbook(filename=os.path.join(tmpdir, filename))
        for sheet in workbook.sheetnames:
            for row in workbook[sheet].iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        assert '"' not in cell.value

def test_task_func_with_non_xlsx_files(tmpdir):
    # Create some non-xlsx files
    with open(os.path.join(tmpdir, 'test.txt'), 'w') as f:
        f.write('This is a text file.')

    with open(os.path.join(tmpdir, 'test.csv'), 'w') as f:
        f.write('This,is,a,csv,file.')

    result = task_func(str(tmpdir))
    assert result == 0

    # Check that no .xlsx files were created or modified
    assert len(glob.glob(os.path.join(tmpdir, '*.xlsx'))) == 0