import pytest
from src_0965 import task_func
import os
from pathlib import Path
import pandas as pd
import docx

def test_task_func_nonexistent_source_directory(tmp_path):
    source_dir = tmp_path / "nonexistent_source"
    target_dir = tmp_path / "target"
    with pytest.raises(FileNotFoundError):
        task_func(str(source_dir), str(target_dir))

def test_task_func_empty_source_directory(tmp_path):
    source_dir = tmp_path / "empty_source"
    source_dir.mkdir()
    target_dir = tmp_path / "target"
    assert task_func(str(source_dir), str(target_dir)) == 0

def test_task_func_single_txt_file(tmp_path):
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    target_dir = tmp_path / "target"
    (source_dir / "test.txt").write_text("Hello\nWorld")
    assert task_func(str(source_dir), str(target_dir)) == 1
    converted_file = target_dir / "test.csv"
    assert converted_file.exists()
    df = pd.read_csv(converted_file)
    assert df.equals(pd.DataFrame({"Text": ["Hello", "World"]}))

def test_task_func_single_docx_file(tmp_path):
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    target_dir = tmp_path / "target"
    doc = docx.Document()
    doc.add_paragraph("Hello")
    doc.add_paragraph("World")
    doc.save(source_dir / "test.docx")
    assert task_func(str(source_dir), str(target_dir)) == 1
    converted_file = target_dir / "test.csv"
    assert converted_file.exists()
    df = pd.read_csv(converted_file)
    assert df.equals(pd.DataFrame({"Text": ["Hello", "World"]}))

def test_task_func_single_csv_file(tmp_path):
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    target_dir = tmp_path / "target"
    df = pd.DataFrame({"Text": ["Hello", "World"]})
    df.to_csv(source_dir / "test.csv", index=False)
    assert task_func(str(source_dir), str(target_dir)) == 1
    converted_file = target_dir / "test.csv"
    assert converted_file.exists()
    df_converted = pd.read_csv(converted_file)
    assert df_converted.equals(df)

def test_task_func_single_xlsx_file(tmp_path):
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    target_dir = tmp_path / "target"
    df = pd.DataFrame({"Text": ["Hello", "World"]})
    df.to_excel(source_dir / "test.xlsx", index=False, engine="openpyxl")
    assert task_func(str(source_dir), str(target_dir)) == 1
    converted_file = target_dir / "test.csv"
    assert converted_file.exists()
    df_converted = pd.read_csv(converted_file)
    assert df_converted.equals(df)

def test_task_func_multiple_files(tmp_path):
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    target_dir = tmp_path / "target"
    (source_dir / "test1.txt").write_text("Hello\nWorld")
    doc = docx.Document()
    doc.add_paragraph("Hello")
    doc.add_paragraph("World")
    doc.save(source_dir / "test2.docx")
    df = pd.DataFrame({"Text": ["Hello", "World"]})
    df.to_csv(source_dir / "test3.csv", index=False)
    df.to_excel(source_dir / "test4.xlsx", index=False, engine="openpyxl")
    assert task_func(str(source_dir), str(target_dir)) == 4
    for i in range(1, 5):
        converted_file = target_dir / f"test{i}.csv"
        assert converted_file.exists()