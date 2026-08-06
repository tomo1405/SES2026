import pytest
from src_0965 import task_func
import os
from pathlib import Path
import pandas as pd
import docx

@pytest.fixture
def setup_directories(tmpdir):
    source_dir = tmpdir.mkdir("source")
    target_dir = tmpdir.mkdir("target")

    # Create some test files in the source directory
    txt_file = source_dir.join("test.txt")
    txt_file.write("Line 1\nLine 2")

    docx_file = source_dir.join("test.docx")
    doc = docx.Document()
    doc.add_paragraph("Paragraph 1")
    doc.add_paragraph("Paragraph 2")
    doc.save(str(docx_file))

    xlsx_file = source_dir.join("test.xlsx")
    df = pd.DataFrame({"Column1": [1, 2], "Column2": ["A", "B"]})
    df.to_excel(xlsx_file, index=False)

    csv_file = source_dir.join("test.csv")
    df.to_csv(csv_file, index=False)

    return str(source_dir), str(target_dir)

def test_task_func(setup_directories):
    source_dir, target_dir = setup_directories

    # Run the function
    converted_files = task_func(source_dir, target_dir)

    # Check that the correct number of files were converted
    assert converted_files == 4

    # Check that the target directory contains the converted files
    target_files = os.listdir(target_dir)
    assert len(target_files) == 4
    assert all(file.endswith(".csv") for file in target_files)

    # Check the contents of each converted file
    for file in target_files:
        df = pd.read_csv(os.path.join(target_dir, file))
        if file.startswith("test_txt"):
            assert df.equals(pd.DataFrame({"Text": ["Line 1", "Line 2"]}))
        elif file.startswith("test_docx"):
            assert df.equals(pd.DataFrame({"Text": ["Paragraph 1", "Paragraph 2"]}))
        elif file.startswith("test_xlsx"):
            assert df.equals(pd.DataFrame({"Column1": [1, 2], "Column2": ["A", "B"]}))
        elif file.startswith("test_csv"):
            assert df.equals(pd.DataFrame({"Column1": [1, 2], "Column2": ["A", "B"]}))

def test_task_func_nonexistent_source_directory():
    with pytest.raises(FileNotFoundError):
        task_func("/nonexistent/source", "/target")

def test_task_func_nonexistent_target_directory(tmpdir):
    source_dir = tmpdir.mkdir("source")
    target_dir = "/nonexistent/target"

    # Create a test file in the source directory
    txt_file = source_dir.join("test.txt")
    txt_file.write("Line 1\nLine 2")

    # Run the function
    converted_files = task_func(str(source_dir), target_dir)

    # Check that the target directory was created and the file was converted
    assert os.path.exists(target_dir)
    assert converted_files == 1
    assert os.path.exists(os.path.join(target_dir, "test_txt.csv"))