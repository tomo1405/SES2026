import pytest
from src_0373 import task_func
import os
import tempfile
from docx import Document

def create_temp_docx(content):
    with tempfile.NamedTemporaryFile(delete=False, suffix='.docx') as temp_file:
        doc = Document()
        doc.add_paragraph(content)
        doc.save(temp_file.name)
        return temp_file.name

def remove_temp_file(file_path):
    os.remove(file_path)

def test_task_func_no_files():
    with tempfile.TemporaryDirectory() as temp_dir:
        result = task_func(temp_dir)
        assert result == 0

def test_task_func_one_file():
    with tempfile.TemporaryDirectory() as temp_dir:
        file_path = create_temp_docx('This is a "test" file.')
        result = task_func(temp_dir)
        assert result == 1
        document = Document(file_path)
        assert document.paragraphs[0].text == 'This is a \"test\" file.'
        remove_temp_file(file_path)

def test_task_func_multiple_files():
    with tempfile.TemporaryDirectory() as temp_dir:
        file1_path = create_temp_docx('First "file".')
        file2_path = create_temp_docx('Second "file".')
        result = task_func(temp_dir)
        assert result == 2
        document1 = Document(file1_path)
        document2 = Document(file2_path)
        assert document1.paragraphs[0].text == 'First \"file\".'
        assert document2.paragraphs[0].text == 'Second \"file\".'
        remove_temp_file(file1_path)
        remove_temp_file(file2_path)

def test_task_func_no_quotes():
    with tempfile.TemporaryDirectory() as temp_dir:
        file_path = create_temp_docx('No quotes here.')
        result = task_func(temp_dir)
        assert result == 1
        document = Document(file_path)
        assert document.paragraphs[0].text == 'No quotes here.'
        remove_temp_file(file_path)

def test_task_func_escaped_quotes():
    with tempfile.TemporaryDirectory() as temp_dir:
        file_path = create_temp_docx('Escaped \"quotes\".')
        result = task_func(temp_dir)
        assert result == 1
        document = Document(file_path)
        assert document.paragraphs[0].text == 'Escaped \"quotes\".'
        remove_temp_file(file_path)