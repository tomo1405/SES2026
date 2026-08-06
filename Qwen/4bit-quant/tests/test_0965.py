import pytest
from src_0965 import task_func
import os
import shutil
import tempfile

def test_task_func_with_non_existent_source_directory():
    with tempfile.TemporaryDirectory() as temp_dir:
        target_dir = os.path.join(temp_dir, "target")
        with pytest.raises(FileNotFoundError):
            task_func("non_existent_source", target_dir)

def test_task_func_with_empty_source_directory():
    with tempfile.TemporaryDirectory() as temp_dir:
        source_dir = os.path.join(temp_dir, "source")
        target_dir = os.path.join(temp_dir, "target")
        assert task_func(source_dir, target_dir) == 0

def test_task_func_with_single_txt_file():
    with tempfile.TemporaryDirectory() as temp_dir:
        source_dir = os.path.join(temp_dir, "source")
        target_dir = os.path.join(temp_dir, "target")
        os.makedirs(source_dir, exist_ok=True)
        with open(os.path.join(source_dir, "test.txt"), "w") as f:
            f.write("Hello\nWorld")

        assert task_func(source_dir, target_dir) == 1
        assert os.path.exists(os.path.join(target_dir, "test.csv"))

def test_task_func_with_multiple_files_of_different_types():
    with tempfile.TemporaryDirectory() as temp_dir:
        source_dir = os.path.join(temp_dir, "source")
        target_dir = os.path.join(temp_dir, "target")
        os.makedirs(source_dir, exist_ok=True)

        # Create a .txt file
        with open(os.path.join(source_dir, "test.txt"), "w") as f:
            f.write("Hello\nWorld")

        # Create a .csv file
        with open(os.path.join(source_dir, "test.csv"), "w") as f:
            f.write("Name,Age\nJohn,30")

        # Create a .xlsx file
        import pandas as pd
        df = pd.DataFrame({"Name": ["John"], "Age": [30]})
        df.to_excel(os.path.join(source_dir, "test.xlsx"), index=False)

        # Create a .docx file
        from docx import Document
        doc = Document()
        doc.add_paragraph("Hello")
        doc.add_paragraph("World")
        doc.save(os.path.join(source_dir, "test.docx"))

        assert task_func(source_dir, target_dir) == 4
        assert os.path.exists(os.path.join(target_dir, "test.csv"))
        assert os.path.exists(os.path.join(target_dir, "test.csv"))
        assert os.path.exists(os.path.join(target_dir, "test.csv"))
        assert os.path.exists(os.path.join(target_dir, "test.csv"))

def test_task_func_with_subdirectories():
    with tempfile.TemporaryDirectory() as temp_dir:
        source_dir = os.path.join(temp_dir, "source")
        target_dir = os.path.join(temp_dir, "target")
        os.makedirs(os.path.join(source_dir, "subdir"), exist_ok=True)

        # Create a .txt file in the subdirectory
        with open(os.path.join(source_dir, "subdir", "test.txt"), "w") as f:
            f.write("Hello\nWorld")

        assert task_func(source_dir, target_dir) == 1
        assert os.path.exists(os.path.join(target_dir, "test.csv"))

def test_task_func_with_existing_target_directory():
    with tempfile.TemporaryDirectory() as temp_dir:
        source_dir = os.path.join(temp_dir, "source")
        target_dir = os.path.join(temp_dir, "target")
        os.makedirs(source_dir, exist_ok=True)
        os.makedirs(target_dir, exist_ok=True)

        with open(os.path.join(source_dir, "test.txt"), "w") as f:
            f.write("Hello\nWorld")

        assert task_func(source_dir, target_dir) == 1
        assert os.path.exists(os.path.join(target_dir, "test.csv"))