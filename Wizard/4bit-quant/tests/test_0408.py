python
import os
import csv
from openpyxl import load_workbook
import pytest

def task_func(file_name, excel_file_path, csv_file_path) -> str:

    excel_file = os.path.join(excel_file_path, file_name)
    # Check if the Excel file exists
    if not os.path.isfile(excel_file):
        raise FileNotFoundError(f"[Errno 2] No such file or directory: '{excel_file}'")

    workbook = load_workbook(filename=excel_file, read_only=True)
    sheet = workbook.active

    data = [[cell.value for cell in row] for row in sheet.iter_rows()]

    csv_file_name = os.path.splitext(file_name)[0] + '.csv'
    csv_file = os.path.join(csv_file_path, csv_file_name)

    with open(csv_file, 'w', newline='', encoding='utf-8') as file:
        writer = csv.writer(file)
        writer.writerows(data)

    return csv_file_name

def test_task_func():
    # Test case 1: Valid input file name, valid Excel file path, valid CSV file path
    file_name = 'test.xlsx'
    excel_file_path = 'C:/Users/user/Documents/test'
    csv_file_path = 'C:/Users/user/Documents/test'
    expected_output = 'test.csv'
    assert task_func(file_name, excel_file_path, csv_file_path) == expected_output

    # Test case 2: Invalid input file name, valid Excel file path, valid CSV file path
    file_name = 'test.txt'
    excel_file_path = 'C:/Users/user/Documents/test'
    csv_file_path = 'C:/Users/user/Documents/test'
    with pytest.raises(FileNotFoundError):
        task_func(file_name, excel_file_path, csv_file_path)

    # Test case 3: Valid input file name, invalid Excel file path, valid CSV file path
    file_name = 'test.xlsx'
    excel_file_path = 'C:/Users/user/Documents/test1'
    csv_file_path = 'C:/Users/user/Documents/test'
    with pytest.raises(FileNotFoundError):
        task_func(file_name, excel_file_path, csv_file_path)

    # Test case 4: Valid input file name, valid Excel file path, invalid CSV file path
    file_name = 'test.xlsx'
    excel_file_path = 'C:/Users/user/Documents/test'
    csv_file_path = 'C:/Users/user/Documents/test1'
    with pytest.raises(FileNotFoundError):
        task_func(file_name, excel_file_path, csv_file_path)